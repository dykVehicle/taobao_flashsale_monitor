# -*- coding: utf-8 -*-
"""
淘宝闪购商品监控工具 - 数据导出模块
"""

import os
import csv
import json
import logging
from typing import List, Dict
from datetime import datetime

logger = logging.getLogger(__name__)


class GoodsExporter:
    """商品数据导出器"""
    
    def __init__(self, output_dir: str = "./exports"):
        """
        初始化导出器
        
        Args:
            output_dir: 导出文件目录
        """
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)
    
    def export_to_excel(self, goods_list: List, store_name: str, 
                        statistics: Dict = None, 
                        problematic_goods: Dict = None) -> str:
        """
        导出到Excel文件
        
        Args:
            goods_list: 商品列表
            store_name: 门店名称
            statistics: 统计信息
            problematic_goods: 问题商品
            
        Returns:
            导出文件路径
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            logger.error("请安装openpyxl: pip install openpyxl")
            raise
        
        date_str = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{store_name}_{date_str}_商品状态报告.xlsx"
        filepath = os.path.join(self.output_dir, filename)
        
        wb = openpyxl.Workbook()
        
        # ====== Sheet1: 统计概览 ======
        ws_summary = wb.active
        ws_summary.title = "统计概览"
        
        # 标题样式
        title_font = Font(bold=True, size=14, color="FFFFFF")
        title_fill = PatternFill(start_color="FF6200", end_color="FF6200", fill_type="solid")
        header_font = Font(bold=True, size=11)
        header_fill = PatternFill(start_color="FFF2E6", end_color="FFF2E6", fill_type="solid")
        
        # 写入标题
        ws_summary.merge_cells('A1:C1')
        ws_summary['A1'] = f"【{store_name}】商品状态报告"
        ws_summary['A1'].font = title_font
        ws_summary['A1'].fill = title_fill
        ws_summary['A1'].alignment = Alignment(horizontal='center')
        
        ws_summary['A2'] = f"统计时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        
        # 写入统计数据
        if statistics:
            ws_summary['A4'] = "商品统计"
            ws_summary['A4'].font = header_font
            
            stats_data = [
                ("总商品数", statistics.get('total', 0)),
                ("在售", statistics.get('on_sale', 0)),
                ("已下架", statistics.get('off_sale', 0)),
                ("已售罄", statistics.get('sold_out', 0)),
                ("缺货", statistics.get('out_of_stock', 0)),
                ("暂停售卖", statistics.get('pause', 0)),
            ]
            
            for i, (label, value) in enumerate(stats_data, 5):
                ws_summary[f'A{i}'] = label
                ws_summary[f'B{i}'] = value
        
        # ====== Sheet2: 全部商品 ======
        ws_all = wb.create_sheet("全部商品")
        headers = ["序号", "商品ID", "商品名称", "分类", "价格", "原价", "库存", "状态"]
        
        for col, header in enumerate(headers, 1):
            cell = ws_all.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
        
        for row, goods in enumerate(goods_list, 2):
            ws_all.cell(row=row, column=1, value=row-1)
            ws_all.cell(row=row, column=2, value=getattr(goods, 'goods_id', ''))
            ws_all.cell(row=row, column=3, value=getattr(goods, 'goods_name', ''))
            ws_all.cell(row=row, column=4, value=getattr(goods, 'category', ''))
            ws_all.cell(row=row, column=5, value=getattr(goods, 'price', 0))
            ws_all.cell(row=row, column=6, value=getattr(goods, 'original_price', 0))
            ws_all.cell(row=row, column=7, value=getattr(goods, 'stock', 0))
            ws_all.cell(row=row, column=8, value=getattr(goods, 'status_text', ''))
            
            # 根据状态设置颜色
            status = getattr(goods, 'status', None)
            if status:
                status_value = status.value if hasattr(status, 'value') else str(status)
                if status_value in ["已下架", "OFF_SALE"]:
                    ws_all.cell(row=row, column=8).fill = PatternFill(
                        start_color="FFCDD2", end_color="FFCDD2", fill_type="solid"
                    )
                elif status_value in ["已售罄", "SOLD_OUT"]:
                    ws_all.cell(row=row, column=8).fill = PatternFill(
                        start_color="FFF9C4", end_color="FFF9C4", fill_type="solid"
                    )
        
        # ====== Sheet3-6: 问题商品分类 ======
        if problematic_goods:
            status_sheets = [
                ("已下架商品", problematic_goods.get("已下架", [])),
                ("已售罄商品", problematic_goods.get("已售罄", [])),
                ("缺货商品", problematic_goods.get("缺货", [])),
                ("暂停售卖商品", problematic_goods.get("暂停售卖", [])),
            ]
            
            for sheet_name, sheet_goods in status_sheets:
                if sheet_goods:
                    ws = wb.create_sheet(sheet_name)
                    
                    for col, header in enumerate(headers, 1):
                        cell = ws.cell(row=1, column=col, value=header)
                        cell.font = header_font
                        cell.fill = header_fill
                    
                    for row, goods in enumerate(sheet_goods, 2):
                        ws.cell(row=row, column=1, value=row-1)
                        ws.cell(row=row, column=2, value=getattr(goods, 'goods_id', ''))
                        ws.cell(row=row, column=3, value=getattr(goods, 'goods_name', ''))
                        ws.cell(row=row, column=4, value=getattr(goods, 'category', ''))
                        ws.cell(row=row, column=5, value=getattr(goods, 'price', 0))
                        ws.cell(row=row, column=6, value=getattr(goods, 'original_price', 0))
                        ws.cell(row=row, column=7, value=getattr(goods, 'stock', 0))
                        ws.cell(row=row, column=8, value=getattr(goods, 'status_text', ''))
        
        # 调整列宽
        for ws in wb.worksheets:
            for col in range(1, 9):
                ws.column_dimensions[get_column_letter(col)].width = 15
            ws.column_dimensions['C'].width = 40  # 商品名称列加宽
        
        wb.save(filepath)
        logger.info(f"Excel导出成功: {filepath}")
        return filepath
    
    def export_to_csv(self, goods_list: List, store_name: str) -> str:
        """
        导出到CSV文件
        
        Args:
            goods_list: 商品列表
            store_name: 门店名称
            
        Returns:
            导出文件路径
        """
        date_str = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{store_name}_{date_str}_商品列表.csv"
        filepath = os.path.join(self.output_dir, filename)
        
        headers = ["序号", "商品ID", "商品名称", "分类", "价格", "原价", "库存", "状态"]
        
        with open(filepath, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            
            for i, goods in enumerate(goods_list, 1):
                writer.writerow([
                    i,
                    getattr(goods, 'goods_id', ''),
                    getattr(goods, 'goods_name', ''),
                    getattr(goods, 'category', ''),
                    getattr(goods, 'price', 0),
                    getattr(goods, 'original_price', 0),
                    getattr(goods, 'stock', 0),
                    getattr(goods, 'status_text', ''),
                ])
        
        logger.info(f"CSV导出成功: {filepath}")
        return filepath
    
    def export_to_json(self, goods_list: List, store_name: str, 
                       statistics: Dict = None) -> str:
        """
        导出到JSON文件
        
        Args:
            goods_list: 商品列表
            store_name: 门店名称
            statistics: 统计信息
            
        Returns:
            导出文件路径
        """
        date_str = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{store_name}_{date_str}_商品数据.json"
        filepath = os.path.join(self.output_dir, filename)
        
        data = {
            "store_name": store_name,
            "export_time": datetime.now().isoformat(),
            "statistics": statistics or {},
            "goods_list": [
                goods.to_dict() if hasattr(goods, 'to_dict') else goods
                for goods in goods_list
            ]
        }
        
        with open(filepath, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        
        logger.info(f"JSON导出成功: {filepath}")
        return filepath

