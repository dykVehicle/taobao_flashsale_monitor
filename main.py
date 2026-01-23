#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
淘宝闪购商品监控工具 - 主程序

功能:
1. 通过Selenium抓取淘宝闪购商家版的商品数据
2. 统计已下架和已售罄的产品
3. 导出商品列表到Excel
4. 通过企业微信机器人发送到门店群

使用方法:
    python main.py                    # 默认持续模式（循环监控）
    python main.py --once             # 仅运行一次检查后退出
    python main.py --export-only      # 仅导出不发送（默认持续模式）
"""

import os
import sys
import argparse
import logging
import json
import subprocess
import time

def _run_silent(cmd, allow_fail=False):
    """静默运行命令"""
    try:
        subprocess.check_call(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        return True
    except (subprocess.CalledProcessError, FileNotFoundError):
        if not allow_fail:
            return False
        return False

def _ensure_pip():
    """确保 pip 已安装"""
    # 检查 pip 是否可用
    result = subprocess.run(
        [sys.executable, "-m", "pip", "--version"],
        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL
    )
    if result.returncode == 0:
        return True
    
    # pip 不存在，尝试安装
    print("正在配置 pip...")
    
    # 方法1: 使用 ensurepip（Python 内置）
    if _run_silent([sys.executable, "-m", "ensurepip", "--upgrade"], allow_fail=True):
        return True
    
    # 方法2: 使用 apt 安装（Debian/Ubuntu）
    if os.path.exists("/usr/bin/apt"):
        if _run_silent(["sudo", "apt", "update", "-qq"], allow_fail=True):
            if _run_silent(["sudo", "apt", "install", "-y", "-qq", "python3-pip"], allow_fail=True):
                return True
    
    # 方法3: 使用 get-pip.py
    try:
        import urllib.request
        import tempfile
        get_pip_url = "https://bootstrap.pypa.io/get-pip.py"
        with tempfile.NamedTemporaryFile(suffix=".py", delete=False) as f:
            urllib.request.urlretrieve(get_pip_url, f.name)
            if _run_silent([sys.executable, f.name, "--quiet"], allow_fail=True):
                os.unlink(f.name)
                return True
            os.unlink(f.name)
    except Exception:
        pass
    
    return False

def ensure_dependencies():
    """自动检查并安装所有依赖（无感知安装）- 包括 pip"""
    required_packages = {
        'requests': 'requests',
        'bs4': 'beautifulsoup4',
        'lxml': 'lxml',
        'openpyxl': 'openpyxl',
        'schedule': 'schedule',
        'selenium': 'selenium',
    }
    
    # 检查缺失的包
    missing = []
    for import_name, package_name in required_packages.items():
        try:
            __import__(import_name)
        except ImportError:
            missing.append(package_name)
    
    if not missing:
        return  # 所有依赖已安装，静默返回
    
    # 确保 pip 可用
    if not _ensure_pip():
        print("✗ 无法安装 pip，请手动安装后重试")
        print("  Ubuntu/Debian: sudo apt install python3-pip")
        print("  其他系统: https://pip.pypa.io/en/stable/installation/")
        sys.exit(1)
    
    # 安装缺失的 Python 包
    print(f"正在配置依赖: {', '.join(missing)}...")
    try:
        subprocess.check_call(
            [sys.executable, "-m", "pip", "install", "-q", "--disable-pip-version-check"] + missing,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL
        )
        print("✓ 依赖配置完成！")
    except subprocess.CalledProcessError:
        # 静默安装失败，尝试使用 requirements.txt
        req_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "requirements.txt")
        if os.path.exists(req_path):
            try:
                subprocess.check_call(
                    [sys.executable, "-m", "pip", "install", "-q", "--disable-pip-version-check", "-r", req_path],
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL
                )
                print("✓ 依赖配置完成！")
            except subprocess.CalledProcessError as e:
                print(f"✗ 依赖安装失败: {e}")
                print(f"  请手动运行: pip install -r {req_path}")
                sys.exit(1)
        else:
            print(f"✗ 依赖安装失败，请手动安装: pip install {' '.join(missing)}")
            sys.exit(1)

ensure_dependencies()

import requests
from datetime import datetime
from typing import List, Dict
from selenium_fetcher import SeleniumGoodsFetcher

# 配置日志
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[logging.StreamHandler(sys.stdout)]
)
logger = logging.getLogger(__name__)


# ============ 配置 ============
CONFIG = {
    # 饿了么连锁商家后台 chain_id（从 URL 中获取：melody.shop.ele.me/app/chain/{chain_id}/...）
    "chain_id": "99760038",
    # 商家后台基础 URL
    "base_url": "https://melody.shop.ele.me",
    "stores": [
        {
            "name": "测试门店",
            "shop_id": "1303549223",  # 如果有多个门店可以在这里配置
            "webhook": ""  # 企业微信机器人webhook地址，留空则不发送
        }
    ],
    "export_dir": "./exports",
    "browser": {
        # 远程调试端口（Chromium/Chrome）
        "debug_port": 9222,
        # 固定用户数据目录（用于"登录一次永久复用登录态"）
        # 默认放在项目目录下：taobao_flashsale_monitor/chromium_profile
        "user_data_dir": os.path.join(os.path.dirname(os.path.abspath(__file__)), "chromium_profile"),
        # 浏览器可执行文件路径（可选；留空自动寻找 Chromium/Chrome/Edge）
        "browser_path": "",
    },
}


def fetch_problematic_goods(fetcher: SeleniumGoodsFetcher, shop_id: str, dump_debug: bool = False) -> Dict:
    """使用Selenium抓取已下架和已售罄的商品（复用同一个浏览器profile保存登录态）"""
    logger.info(f"开始抓取店铺 {shop_id} 的问题商品...")

    # 复用同一个 fetcher/driver，但切换 shopId
    fetcher.shop_id = shop_id
    goods_list = fetcher.login_and_fetch(
        auto_login=False,
        wait_for_login=True,        # 首次需要你在浏览器里登录；后续会自动复用登录态
        login_timeout=30 * 60,      # 最多等待30分钟，避免扫码/短信流程超时
        dump_debug=dump_debug,
    )
    
    # 分类统计
    off_sale_goods = []
    sold_out_goods = []
    
    for goods in goods_list:
        if goods.status == "OFF_SALE":
            off_sale_goods.append(goods)
        elif goods.status == "SOLD_OUT":
            sold_out_goods.append(goods)
    
    return {
        "off_sale": off_sale_goods,
        "sold_out": sold_out_goods,
        "total": len(goods_list)
    }


def export_to_excel(goods_data: Dict, store_name: str, output_dir: str) -> str:
    """导出到Excel文件"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        logger.error("请安装openpyxl: pip install openpyxl")
        return ""
    
    os.makedirs(output_dir, exist_ok=True)
    
    wb = Workbook()
    
    # 已下架商品表
    ws_off_sale = wb.active
    ws_off_sale.title = "已下架商品"
    ws_off_sale.append(["序号", "商品名称", "价格", "库存", "状态"])
    
    # 设置标题样式
    header_fill = PatternFill(start_color="FF6200", end_color="FF6200", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    for cell in ws_off_sale[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    
    for i, goods in enumerate(goods_data.get("off_sale", []), 1):
        ws_off_sale.append([i, goods.goods_name, goods.price, goods.stock, goods.status_text])
    
    # 已售罄商品表
    ws_sold_out = wb.create_sheet("已售罄商品")
    ws_sold_out.append(["序号", "商品名称", "价格", "库存", "状态"])
    
    for cell in ws_sold_out[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    
    for i, goods in enumerate(goods_data.get("sold_out", []), 1):
        ws_sold_out.append([i, goods.goods_name, goods.price, goods.stock, goods.status_text])
    
    # 调整列宽
    for ws in [ws_off_sale, ws_sold_out]:
        ws.column_dimensions['B'].width = 40
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 12
    
    # 保存
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{store_name}_商品监控报告_{timestamp}.xlsx"
    filepath = os.path.join(output_dir, filename)
    wb.save(filepath)
    
    logger.info(f"已导出到: {filepath}")
    return filepath


def send_wecom_notification(webhook: str, store_name: str, goods_data: Dict):
    """发送企业微信通知"""
    if not webhook:
        logger.info("未配置webhook，跳过发送通知")
        return
    
    off_sale_count = len(goods_data.get("off_sale", []))
    sold_out_count = len(goods_data.get("sold_out", []))
    
    if off_sale_count == 0 and sold_out_count == 0:
        logger.info("无问题商品，跳过发送通知")
        return
    
    # 构建消息内容
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    
    content = f"""## 🔔 【{store_name}】商品监控报告
> 检查时间: {now}

### 📊 统计概览
- **已下架商品**: <font color="warning">{off_sale_count}</font> 个
- **已售罄商品**: <font color="warning">{sold_out_count}</font> 个
"""
    
    # 已下架商品列表
    if off_sale_count > 0:
        content += "\n### 🔴 已下架商品\n"
        for i, goods in enumerate(goods_data["off_sale"][:10], 1):
            content += f"{i}. {goods.goods_name}\n"
        if off_sale_count > 10:
            content += f"... 还有 {off_sale_count - 10} 个\n"
    
    # 已售罄商品列表
    if sold_out_count > 0:
        content += "\n### 🟡 已售罄商品\n"
        for i, goods in enumerate(goods_data["sold_out"][:10], 1):
            content += f"{i}. {goods.goods_name}\n"
        if sold_out_count > 10:
            content += f"... 还有 {sold_out_count - 10} 个\n"
    
    content += "\n> 请及时处理以上商品问题！"
    
    # 发送Markdown消息
    data = {
        "msgtype": "markdown",
        "markdown": {
            "content": content
        }
    }
    
    try:
        response = requests.post(webhook, json=data, timeout=10)
        if response.status_code == 200:
            result = response.json()
            if result.get("errcode") == 0:
                logger.info("企业微信通知发送成功")
            else:
                logger.error(f"发送失败: {result.get('errmsg')}")
        else:
            logger.error(f"发送失败: HTTP {response.status_code}")
    except Exception as e:
        logger.error(f"发送异常: {e}")


def print_report(store_name: str, goods_data: Dict):
    """打印报告"""
    off_sale_count = len(goods_data.get("off_sale", []))
    sold_out_count = len(goods_data.get("sold_out", []))
    
    print("\n" + "=" * 60)
    print(f"【{store_name}】商品监控报告")
    print("=" * 60)
    print(f"检查时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print("-" * 60)
    print(f"已下架商品: {off_sale_count} 个")
    print(f"已售罄商品: {sold_out_count} 个")
    print("-" * 60)
    
    if off_sale_count > 0:
        print("\n🔴 已下架商品:")
        for i, goods in enumerate(goods_data["off_sale"], 1):
            print(f"  {i}. {goods.goods_name}")
    
    if sold_out_count > 0:
        print("\n🟡 已售罄商品:")
        for i, goods in enumerate(goods_data["sold_out"], 1):
            print(f"  {i}. {goods.goods_name}")
    
    print("\n" + "=" * 60)


def run_monitor(fetcher: SeleniumGoodsFetcher, export_only: bool = False, export_always: bool = False, dump_debug: bool = False):
    """运行一次监控（单轮）"""
    for store in CONFIG["stores"]:
        store_name = store["name"]
        shop_id = store["shop_id"]
        webhook = store.get("webhook", "")
        
        logger.info(f"正在检查门店: {store_name}")
        
        try:
            # 抓取数据
            goods_data = fetch_problematic_goods(fetcher, shop_id, dump_debug=dump_debug)
            
            # 打印报告
            print_report(store_name, goods_data)
            
            # 导出Excel（持续模式下避免无限生成空报表：无问题商品默认跳过）
            off_sale_count = len(goods_data.get("off_sale", []))
            sold_out_count = len(goods_data.get("sold_out", []))
            if export_only or export_always or off_sale_count > 0 or sold_out_count > 0:
                export_to_excel(goods_data, store_name, CONFIG["export_dir"])
            else:
                logger.info("无问题商品，跳过导出Excel（如需强制导出可使用 --export-always）")
            
            # 发送通知
            if not export_only:
                send_wecom_notification(webhook, store_name, goods_data)
                
        except Exception as e:
            logger.error(f"检查门店 {store_name} 失败: {e}")
            import traceback
            traceback.print_exc()


def main():
    """主函数"""
    parser = argparse.ArgumentParser(description="淘宝闪购商品监控工具")
    parser.add_argument("--export-only", "-e", action="store_true", help="仅导出不发送通知")
    parser.add_argument("--export-always", action="store_true", help="即使无问题商品也导出Excel（持续模式下不推荐）")
    parser.add_argument("--once", action="store_true", help="仅运行一次后退出（默认持续模式）")
    parser.add_argument("--interval", "-i", type=int, default=30, help="持续模式下每次检查间隔（分钟），默认30")
    parser.add_argument("--debug-port", type=int, default=CONFIG["browser"]["debug_port"], help="浏览器远程调试端口，默认9222")
    parser.add_argument("--profile-dir", type=str, default=CONFIG["browser"]["user_data_dir"], help="浏览器用户数据目录（用于持久化登录态）")
    parser.add_argument("--browser-path", type=str, default=CONFIG["browser"]["browser_path"], help="浏览器可执行文件路径（可选，留空自动寻找）")
    parser.add_argument("--no-open-browser", action="store_true", help="不自动启动浏览器（若你已手动启动远程调试浏览器）")
    parser.add_argument("--dump-debug", action="store_true", help="保存调试截图/HTML（debug_screenshot.png/debug_page.html）")
    args = parser.parse_args()
    
    print("\n" + "=" * 60)
    print("淘宝闪购商品监控工具")
    print("=" * 60)
    print("使用说明:")
    print("1. 程序会自动启动 Chromium/Chrome（远程调试模式）并复用固定 profile 保存登录态")
    print(f"   - 远程调试端口: {args.debug_port}")
    print(f"   - profile目录(勿删除): {args.profile_dir}")
    print("2. 首次运行请在打开的浏览器里登录一次（后续会自动复用登录态）")
    print("3. 默认持续模式：脚本会一直运行，按间隔重复检查（Ctrl+C 退出）")
    print("   如需只跑一次后退出：加参数 --once")
    print("=" * 60 + "\n")
    
    # 初始化 Selenium 抓取器：使用固定 profile 目录来"登录一次永久复用"
    fetcher = SeleniumGoodsFetcher(
        shop_id=CONFIG["stores"][0]["shop_id"],
        chain_id=CONFIG["chain_id"],
        base_url=CONFIG["base_url"],
        headless=False,
        debug_port=args.debug_port,
        user_data_dir=args.profile_dir,
        browser_path=(args.browser_path or None),
        auto_launch_browser=(not args.no_open_browser),
    )
    
    # 启动时先拉起浏览器（更符合"运行脚本后自动打开 chromium"）
    if not args.no_open_browser:
        # 饿了么连锁商家后台首页
        open_url = f"{CONFIG['base_url']}/app/chain/{CONFIG['chain_id']}/shop#app.chainshop.shop"
        fetcher.ensure_debug_browser(open_url=open_url)
    
    try:
        if args.once:
            run_monitor(fetcher, export_only=args.export_only, export_always=args.export_always, dump_debug=args.dump_debug)
        else:
            logger.info(f"已启动持续模式：每 {args.interval} 分钟检查一次（Ctrl+C 退出）")
            while True:
                run_monitor(fetcher, export_only=args.export_only, export_always=args.export_always, dump_debug=args.dump_debug)
                logger.info(f"本轮检查结束，等待 {args.interval} 分钟后进入下一轮...")
                time.sleep(max(1, args.interval) * 60)
    except KeyboardInterrupt:
        logger.info("收到退出信号，正在结束...")
    finally:
        # 仅关闭 webdriver 会话；浏览器 profile 会保留登录态
        try:
            fetcher.close()
        except Exception:
            pass


if __name__ == "__main__":
    main()
