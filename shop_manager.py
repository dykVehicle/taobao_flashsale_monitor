#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
店铺管理模块 - 管理多店铺配置
"""

import os
from dataclasses import dataclass
from typing import List, Optional
import json


@dataclass
class ShopInfo:
    """店铺信息"""
    name: str  # 店铺名称（用于搜索）
    webhook: str  # 企业微信webhook地址
    enabled: bool = True  # 是否启用监控


class ShopManager:
    """店铺管理器"""
    
    def __init__(self, config_path: str = None):
        """
        初始化店铺管理器
        
        Args:
            config_path: 配置文件路径，支持xlsx或json格式
        """
        self.shops: List[ShopInfo] = []
        self.config_path = config_path
        self.valid_goods_names: set = set()  # 有效商品名称白名单（从"导出商品"sheet加载）
        
        if config_path and os.path.exists(config_path):
            self.load(config_path)
    
    def load(self, config_path: str) -> bool:
        """加载店铺配置"""
        try:
            ext = os.path.splitext(config_path)[1].lower()
            
            if ext in ['.xlsx', '.xls']:
                return self._load_excel(config_path)
            elif ext == '.json':
                return self._load_json(config_path)
            else:
                print(f"不支持的配置文件格式: {ext}")
                return False
                
        except Exception as e:
            print(f"加载店铺配置失败: {e}")
            return False
    
    def _load_excel(self, path: str) -> bool:
        """从Excel加载店铺配置"""
        # 优先尝试使用 openpyxl 直接读取（更可靠，不依赖 pandas）
        try:
            # 一次性加载门店列表和商品白名单（避免重复打开文件）
            return self._load_excel_openpyxl(path)
        except Exception as e:
            print(f"openpyxl 读取失败: {e}，尝试使用 pandas...")
        
        # 备用方案：使用 pandas
        try:
            import pandas as pd
            df = pd.read_excel(path)
            
            # 兼容多种列名
            name_col = None
            webhook_col = None
            
            # 查找门店名称列（兼容"店铺名称"和"门店名称"）
            for col in ['店铺名称', '门店名称', '名称', 'name']:
                if col in df.columns:
                    name_col = col
                    break
            
            # 查找WebHook列（兼容大小写）
            for col in ['WebHook', 'webhook', 'Webhook', 'WEBHOOK']:
                if col in df.columns:
                    webhook_col = col
                    break
            
            if not name_col:
                print(f"Excel缺少必要列: 店铺名称/门店名称，当前列: {df.columns.tolist()}")
                return False
            if not webhook_col:
                print(f"Excel缺少必要列: WebHook，当前列: {df.columns.tolist()}")
                return False
            
            self.shops = []
            for _, row in df.iterrows():
                name = str(row[name_col]).strip()
                webhook = str(row[webhook_col]).strip()
                
                # 过滤空行和无效数据（nan、空字符串等）
                if not name or name.lower() == 'nan':
                    continue
                if not webhook or webhook.lower() == 'nan' or not webhook.startswith('http'):
                    continue
                
                enabled = True
                if '启用' in df.columns:
                    enabled = bool(row.get('启用', True))
                
                self.shops.append(ShopInfo(
                    name=name,
                    webhook=webhook,
                    enabled=enabled
                ))
            
            print(f"成功加载 {len(self.shops)} 个店铺配置 (pandas)")
            return True
            
        except ImportError as e:
            print(f"pandas 导入失败: {e}")
            return False
        except Exception as e:
            print(f"pandas 读取Excel失败: {e}")
            return False
    
    def _load_excel_openpyxl(self, path: str) -> bool:
        """使用 openpyxl 直接读取 Excel（不依赖 pandas）
        
        同时加载：
        1. 门店列表（从 Input/active sheet）
        2. 商品白名单（从"导出商品"sheet，如果存在）
        """
        from openpyxl import load_workbook
        
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        
        # ========== 1. 加载门店列表 ==========
        # 读取表头
        headers = []
        for cell in ws[1]:
            headers.append(str(cell.value or '').strip())
        
        # 查找列索引
        name_idx = None
        webhook_idx = None
        
        for i, header in enumerate(headers):
            if header in ['店铺名称', '门店名称', '名称', 'name']:
                name_idx = i
            elif header.lower() in ['webhook']:
                webhook_idx = i
        
        if name_idx is None:
            print(f"Excel缺少必要列: 店铺名称/门店名称，当前列: {headers}")
            wb.close()
            return False
        if webhook_idx is None:
            print(f"Excel缺少必要列: WebHook，当前列: {headers}")
            wb.close()
            return False
        
        self.shops = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if len(row) <= max(name_idx, webhook_idx):
                continue
            
            name = str(row[name_idx] or '').strip()
            webhook = str(row[webhook_idx] or '').strip()
            
            # 过滤空行和无效数据
            if not name or name.lower() == 'nan' or name == 'None':
                continue
            if not webhook or webhook.lower() == 'nan' or not webhook.startswith('http'):
                continue
            
            self.shops.append(ShopInfo(
                name=name,
                webhook=webhook,
                enabled=True
            ))
        
        print(f"✓ 成功加载 {len(self.shops)} 个店铺配置")
        
        # ========== 2. 加载商品白名单（在同一个 workbook 中） ==========
        self.valid_goods_names = set()
        
        # 查找"导出商品"sheet
        goods_sheet_name = None
        for sheet_name in wb.sheetnames:
            if '导出商品' in sheet_name:
                goods_sheet_name = sheet_name
                break
        
        if goods_sheet_name:
            try:
                goods_ws = wb[goods_sheet_name]
                goods_name_col = 3  # 商品名称在第4列（索引3）
                
                for row in goods_ws.iter_rows(min_row=4, values_only=True):
                    if len(row) <= goods_name_col:
                        continue
                    
                    goods_name = row[goods_name_col]
                    if goods_name and str(goods_name).strip() and str(goods_name).lower() != 'nan':
                        self.valid_goods_names.add(str(goods_name).strip())
                
                if self.valid_goods_names:
                    print(f"✓ 已加载 {len(self.valid_goods_names)} 个商品白名单")
            except Exception as e:
                print(f"⚠ 加载商品白名单失败: {e}")
        
        wb.close()
        return True
    
    def _load_json(self, path: str) -> bool:
        """从JSON加载店铺配置"""
        try:
            with open(path, 'r', encoding='utf-8') as f:
                data = json.load(f)
            
            self.shops = []
            for item in data.get('shops', []):
                self.shops.append(ShopInfo(
                    name=item['name'],
                    webhook=item['webhook'],
                    enabled=item.get('enabled', True)
                ))
            
            return True
        except Exception as e:
            print(f"读取JSON失败: {e}")
            return False
    
    def save_json(self, path: str) -> bool:
        """保存到JSON文件"""
        try:
            data = {
                'shops': [
                    {
                        'name': shop.name,
                        'webhook': shop.webhook,
                        'enabled': shop.enabled
                    }
                    for shop in self.shops
                ]
            }
            with open(path, 'w', encoding='utf-8') as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
            return True
        except Exception as e:
            print(f"保存JSON失败: {e}")
            return False
    
    def get_enabled_shops(self) -> List[ShopInfo]:
        """获取启用的店铺列表"""
        return [shop for shop in self.shops if shop.enabled]
    
    def get_shop_by_name(self, name: str) -> Optional[ShopInfo]:
        """根据名称获取店铺"""
        for shop in self.shops:
            if shop.name == name:
                return shop
        return None
    
    def add_shop(self, name: str, webhook: str, enabled: bool = True):
        """添加店铺"""
        self.shops.append(ShopInfo(name=name, webhook=webhook, enabled=enabled))
    
    def remove_shop(self, name: str) -> bool:
        """移除店铺"""
        for i, shop in enumerate(self.shops):
            if shop.name == name:
                self.shops.pop(i)
                return True
        return False
    
    def _load_valid_goods_from_excel(self, path: str) -> bool:
        """
        从Excel的"导出商品"sheet加载有效商品名称白名单
        
        Args:
            path: Excel文件路径
            
        Returns:
            是否成功加载
        """
        try:
            from openpyxl import load_workbook
            
            wb = load_workbook(path, read_only=True, data_only=True)
            
            # 查找"导出商品"sheet
            sheet_name = None
            for name in wb.sheetnames:
                if '导出商品' in name:
                    sheet_name = name
                    break
            
            if not sheet_name:
                print("未找到'导出商品' sheet，跳过商品白名单加载")
                wb.close()
                return False
            
            ws = wb[sheet_name]
            self.valid_goods_names = set()
            
            # 商品名称在第4列（索引3），从第4行开始（跳过前3行表头/说明）
            goods_name_col = 3  # 第4列
            
            for row_idx, row in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
                if len(row) <= goods_name_col:
                    continue
                
                goods_name = row[goods_name_col]
                if goods_name and str(goods_name).strip() and str(goods_name).lower() != 'nan':
                    # 清理商品名称（去除首尾空格）
                    clean_name = str(goods_name).strip()
                    self.valid_goods_names.add(clean_name)
            
            wb.close()
            print(f"✓ 已加载 {len(self.valid_goods_names)} 个有效商品名称（白名单）")
            return True
            
        except Exception as e:
            print(f"加载商品白名单失败: {e}")
            return False
    
    def is_valid_goods(self, goods_name: str) -> bool:
        """
        检查商品名称是否在白名单中
        
        Args:
            goods_name: 商品名称
            
        Returns:
            是否有效（在白名单中）
        """
        if not self.valid_goods_names:
            # 如果没有白名单，所有商品都认为有效
            return True
        
        # 精确匹配
        if goods_name in self.valid_goods_names:
            return True
        
        # 模糊匹配：检查是否包含（处理名称可能有差异的情况）
        clean_name = goods_name.strip()
        for valid_name in self.valid_goods_names:
            # 名称相互包含则认为匹配
            if clean_name in valid_name or valid_name in clean_name:
                return True
        
        return False
    
    def filter_valid_goods(self, goods_list: list) -> tuple:
        """
        过滤商品列表，只保留在白名单中的商品
        
        Args:
            goods_list: 商品列表（GoodsItem对象列表）
            
        Returns:
            (有效商品列表, 无效商品列表)
        """
        if not self.valid_goods_names:
            # 如果没有白名单，所有商品都认为有效
            return goods_list, []
        
        valid_goods = []
        invalid_goods = []
        
        for goods in goods_list:
            goods_name = getattr(goods, 'goods_name', '') or getattr(goods, 'name', '')
            if self.is_valid_goods(goods_name):
                valid_goods.append(goods)
            else:
                invalid_goods.append(goods)
        
        return valid_goods, invalid_goods
    
    def get_whitelist_count(self) -> int:
        """获取白名单中的商品数量"""
        return len(self.valid_goods_names)


# 测试代码
if __name__ == "__main__":
    manager = ShopManager("doc/店铺列表.xlsx")
    for shop in manager.shops:
        print(f"店铺: {shop.name}, Webhook: {shop.webhook[:50]}...")
